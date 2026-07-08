"""
Excel 生成工具 — 基于 openpyxl。

支持：多 Sheet、表头加粗、自动列宽、暗色主题。
"""

import logging
import os
from pathlib import Path

logger = logging.getLogger("nano_agent.tools.excel")


def _ensure_openpyxl():
    """确保 openpyxl 已安装。"""
    try:
        import openpyxl  # noqa: F401
        return
    except ImportError:
        raise ImportError(
            "openpyxl is not installed. Run: pip install openpyxl"
        )


class Excel:
    TOOLS = [
        ("create_excel",
         "Generate an Excel spreadsheet (.xlsx) from structured data. "
         "Supports multiple sheets, table headers, and auto column width. "
         "Returns a download link for the generated file.",
         "create_excel",
         {"sheets": {"type": "array",
                     "description": "List of sheet objects. Each sheet has: name (sheet name), "
                     "headers (list of column headers), rows (list of lists, each inner list is one data row).",
                     "items": {"type": "object",
                               "properties": {
                                   "name": {"type": "string", "description": "Sheet/tab name"},
                                   "headers": {"type": "array", "description": "Column header names",
                                               "items": {"type": "string"}},
                                   "rows": {"type": "array", "description": "Data rows",
                                            "items": {"type": "array", "items": {"type": "string"}}},
                               },
                               "required": ["name", "headers", "rows"]}},
          "filename": {"type": "string", "description": "Output filename without extension (optional, defaults to 'data')"}},
         ["sheets"]),
    ]

    def __init__(self, work_dir: str, charts_dir: str = ""):
        self.work_dir = work_dir
        self.charts_dir = charts_dir or work_dir

    def create_excel(self, sheets: list[dict], filename: str = "data") -> str:
        """生成 Excel 文件。

        Args:
            sheets: 工作表列表。每项:
                    - name: sheet 名称
                    - headers: 列标题列表
                    - rows: 数据行列表
            filename: 输出文件名（不含路径和扩展名）

        Returns:
            生成结果，包含下载链接。
        """
        _ensure_openpyxl()
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        # 删除默认的 Sheet，后面自己创建
        wb.remove(wb.active)

        # 样式
        header_font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="7C3AED", end_color="7C3AED", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_font = Font(name="Microsoft YaHei", size=10, color="E0E0E0")
        cell_fill = PatternFill(start_color="1E1E2E", end_color="1E1E2E", fill_type="solid")
        cell_align = Alignment(vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin", color="333333"),
            right=Side(style="thin", color="333333"),
            top=Side(style="thin", color="333333"),
            bottom=Side(style="thin", color="333333"),
        )

        for sheet_data in sheets:
            name = sheet_data.get("name", "Sheet")
            headers = sheet_data.get("headers", [])
            rows = sheet_data.get("rows", [])

            # sheet 名最长 31 字符
            ws = wb.create_sheet(title=name[:31])

            # 写表头
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=str(header))
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align
                cell.border = thin_border

            # 写数据行
            for row_idx, row_data in enumerate(rows, 2):
                for col_idx, value in enumerate(row_data, 1):
                    cell = ws.cell(row=row_idx, column=col_idx,
                                   value=str(value) if value is not None else "")
                    cell.font = cell_font
                    cell.fill = cell_fill
                    cell.alignment = cell_align
                    cell.border = thin_border

            # 自动列宽
            for col_idx in range(1, len(headers) + 1):
                max_width = len(str(headers[col_idx - 1])) * 2  # 中文宽字符
                for row_idx in range(2, len(rows) + 2):
                    val = ws.cell(row=row_idx, column=col_idx).value
                    if val:
                        # 粗略估算宽度：中文字符算2，其他算1
                        char_width = sum(2 if ord(c) > 127 else 1 for c in str(val))
                        max_width = max(max_width, char_width)
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max_width + 4, 60)

            # 冻结首行
            ws.freeze_panes = "A2"

        import re as _re, time as _t
        safe_name = _re.sub(r'[^a-zA-Z0-9._-]', '_', filename)
        safe_name = safe_name.strip("_") or "data"
        safe_name = f"{safe_name}.xlsx"
        # 兜底：如果 sanitize 后仍然只有下划线/点，用时间戳命名
        if not safe_name.replace(".xlsx", "").strip("_"):
            safe_name = f"data_{int(_t.time())}.xlsx"

        # 输出到 charts_dir 以便 Web 下载
        import os as _os
        os.makedirs(self.charts_dir, exist_ok=True)
        filepath = _os.path.join(self.charts_dir, safe_name)
        wb.save(filepath)

        sheet_count = len(sheets)
        total_rows = sum(len(s.get("rows", [])) for s in sheets)
        url = f"/charts/{safe_name}"
        return (
            f"Excel generated: {safe_name}\n"
            f"{sheet_count} sheet(s), {total_rows} data rows\n"
            f"[Download {safe_name}]({url})"
        )
